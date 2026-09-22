from pathlib import Path

import pandas as pd
import openpyxl
from openpyxl.styles import Font


# ============================================================
# PATH CONFIGURATION
# ============================================================

SCRIPT_DIR = Path(__file__).resolve().parent
PROJECT_ROOT = SCRIPT_DIR.parents[3]

DOUBLE_CATEGORIZATION_DIR = (
    SCRIPT_DIR.parent / "Double_Categorization"
)

MODEL_FILE = DOUBLE_CATEGORIZATION_DIR / "MultiLabel.xlsx"

MANUAL_FILE = PROJECT_ROOT / "data" / "categorie_manuali_doppie.csv"

OUTPUT_FILE = DOUBLE_CATEGORIZATION_DIR / "MultiLabel_Evaluation.xlsx"

# ============================================================
# DATA LOADING
# ============================================================

print("Loading files...")

df_modello = pd.read_excel(MODEL_FILE)

try:
    df_manuale = pd.read_csv(
        MANUAL_FILE,
        sep=None,
        engine="python",
        on_bad_lines="skip",
        encoding="utf-8-sig"
    )
    print(f"Successfully loaded '{MANUAL_FILE.name}'.")

except Exception as error:
    raise RuntimeError(
        f"Unable to load the manual annotation file: {MANUAL_FILE}"
    ) from error


# Remove accidental whitespace from column names.
df_modello.columns = df_modello.columns.str.strip()
df_manuale.columns = df_manuale.columns.str.strip()


# ============================================================
# DATA PREPARATION
# ============================================================

# Create a normalized merge key from the evidence text.
df_modello["EVIDENZA_CLEAN"] = (
    df_modello["EVIDENZA"]
    .astype(str)
    .str.strip()
    .str.lower()
)

df_manuale["EVIDENZA_CLEAN"] = (
    df_manuale["EVIDENZA"]
    .astype(str)
    .str.strip()
    .str.lower()
)


# Canonical category labels used throughout the experiment.
CANONICAL_CATEGORIES = {
    "ATTORI / SOGGETTI",
    "TATTICHE / METODO",
    "MEZZI / STRUMENTI",
    "OBIETTIVO / TARGET",
    "CONTENUTO / NARRATIVA",
    "IMPATTO / EFFETTO",
    "VALIDAZIONE ANALITICA",
    "CONTESTO / ECOSISTEMA",
    "FINANZIAMENTO / MODELLO ECONOMICO",
}


def normalize_category(value):
    """Normalize spacing and case and map a label to a canonical category."""

    if pd.isna(value):
        return None

    label = str(value).strip().upper()

    if not label or label == "NAN":
        return None

    # Collapse repeated whitespace and normalize spacing around '/'.
    label = " ".join(label.split())
    label = " / ".join(part.strip() for part in label.split("/"))

    if label not in CANONICAL_CATEGORIES:
        raise ValueError(
            f"Unknown category after normalization: "
            f"{value!r} -> {label!r}"
        )

    return label


# ============================================================
# GROUND-TRUTH PREPARATION
# ============================================================

# Build the ground-truth set from the first and optional
# second manually assigned category.
if (
    "CATEGORIA_MANUALE_1" in df_manuale.columns
    and "CATEGORIA_MANUALE_2" in df_manuale.columns
):

    def create_gt_set(row):
        c1 = normalize_category(row["CATEGORIA_MANUALE_1"])
        c2 = normalize_category(row["CATEGORIA_MANUALE_2"])

        if c1 is None:
            raise ValueError("CATEGORIA_MANUALE_1 cannot be empty.")

        gt_set = {c1} if c2 is None else {c1, c2}

        return gt_set, len(gt_set)


    gt_results = df_manuale.apply(create_gt_set, axis=1)

    df_manuale["GT_SET"] = [r[0] for r in gt_results]
    df_manuale["NUM_GT"] = [r[1] for r in gt_results]

else:
    raise ValueError(
        "The manual file must contain 'CATEGORIA_MANUALE_1' "
        "and 'CATEGORIA_MANUALE_2'."
    )


# ============================================================
# DATA MERGE
# ============================================================

df_merged = df_modello.merge(
    df_manuale[
        [
            "EVIDENZA_CLEAN",
            "GT_SET",
            "NUM_GT",
        ]
    ],
    on="EVIDENZA_CLEAN",
    how="inner",
)


# ============================================================
# EVALUATION
# ============================================================

if not df_merged.empty:

    def evaluate_row(row):
        gt_set = row["GT_SET"]
        num_gt = row["NUM_GT"]

        p1_raw = str(row["CATEGORIA_PRINCIPALE"]).strip().upper()
        p2_raw = str(row["CATEGORIA_SECONDARIA"]).strip().upper()

        if p1_raw in ["ERRORE_FORMATO", "ERRORE_CRITICO"]:
            return "MODEL_ERROR"

        p1 = normalize_category(row["CATEGORIA_PRINCIPALE"])
        p2 = normalize_category(row["CATEGORIA_SECONDARIA"])

        model_set = {
            label
            for label in (p1, p2)
            if label is not None
        }

        # Case A: single-label ground truth.
        if num_gt == 1:
            return (
                "MATCH_SINGLE_LABEL"
                if gt_set.intersection(model_set)
                else "NO_MATCH"
            )

        # Case B: multi-label ground truth.
        num_match = len(gt_set.intersection(model_set))

        if num_match == 2:
            return "MATCH_MULTI_COMPLETE"

        elif num_match == 1:
            return "MATCH_MULTI_PARTIAL"

        return "NO_MATCH"


    df_merged["OUTCOME"] = df_merged.apply(
        evaluate_row,
        axis=1,
    )


    # ========================================================
    # EVIDENCE-LEVEL METRICS
    # ========================================================

    totale = len(df_merged)

    totale_singoli = len(
        df_merged[df_merged["NUM_GT"] == 1]
    )

    totale_doppi = len(
        df_merged[df_merged["NUM_GT"] == 2]
    )

    match_singolo = len(
        df_merged[
            df_merged["OUTCOME"] == "MATCH_SINGLE_LABEL"
        ]
    )

    match_perfetto = len(
        df_merged[
            df_merged["OUTCOME"] == "MATCH_MULTI_COMPLETE"
        ]
    )

    match_parziale = len(
        df_merged[
            df_merged["OUTCOME"] == "MATCH_MULTI_PARTIAL"
        ]
    )

    nessun_match = len(
        df_merged[
            df_merged["OUTCOME"] == "NO_MATCH"
        ]
    )

    errori_modello = len(
        df_merged[
            df_merged["OUTCOME"] == "MODEL_ERROR"
        ]
    )

    match_rate_singolo = (
        match_singolo / totale_singoli
        if totale_singoli
        else 0.0
    )

    complete_match_rate_multi = (
        match_perfetto / totale_doppi
        if totale_doppi
        else 0.0
    )

    partial_match_rate_multi = (
        match_parziale / totale_doppi
        if totale_doppi
        else 0.0
    )

    any_match_multi = (
        (match_perfetto + match_parziale) / totale_doppi
        if totale_doppi
        else 0.0
    )

    evidence_any_match = (
        (
            match_singolo
            + match_perfetto
            + match_parziale
        )
        / totale
        if totale
        else 0.0
    )


    # ========================================================
    # LABEL-LEVEL METRICS
    # ========================================================

    # Micro-averaged multi-label metrics:
    # TP = correctly predicted ground-truth labels
    # FP = predicted labels not present in the ground truth
    # FN = ground-truth labels missed by the model
    tp = fp = fn = 0

    for _, row in df_merged.iterrows():

        if row["OUTCOME"] == "MODEL_ERROR":
            continue

        gt_set = set(row["GT_SET"])

        p1 = normalize_category(
            row["CATEGORIA_PRINCIPALE"]
        )

        p2 = normalize_category(
            row["CATEGORIA_SECONDARIA"]
        )

        pred_set = {
            label
            for label in (p1, p2)
            if label is not None
        }

        tp += len(gt_set & pred_set)
        fp += len(pred_set - gt_set)
        fn += len(gt_set - pred_set)

    micro_precision = (
        tp / (tp + fp)
        if (tp + fp)
        else 0.0
    )

    micro_recall = (
        tp / (tp + fn)
        if (tp + fn)
        else 0.0
    )

    micro_f1 = (
        2
        * micro_precision
        * micro_recall
        / (micro_precision + micro_recall)
        if (micro_precision + micro_recall)
        else 0.0
    )


    # ========================================================
    # CONSOLE REPORT
    # ========================================================

    print("\n" + "=" * 72)
    print("       MULTI-LABEL GROUND-TRUTH EVALUATION REPORT")
    print("=" * 72)

    print(f"Total evidence items:                     {totale}")
    print(f"  - Single-label evidence items:          {totale_singoli}")
    print(f"  - Multi-label evidence items:           {totale_doppi}")

    print("-" * 72)

    print(f"Single-label match (1/1):                 {match_singolo}")
    print(f"Complete multi-label match (2/2):         {match_perfetto}")
    print(f"Partial multi-label match (1/2):          {match_parziale}")
    print(f"No match:                                 {nessun_match}")
    print(f"Model errors:                             {errori_modello}")

    print("-" * 72)
    print("EVIDENCE-LEVEL METRICS")

    print(
        f"Single-label Match Rate:                  "
        f"{match_rate_singolo:.2%}"
    )
    print(
        f"Complete Multi-label Match Rate (2/2):    "
        f"{complete_match_rate_multi:.2%}"
    )
    print(
        f"Partial Multi-label Match Rate (1/2):     "
        f"{partial_match_rate_multi:.2%}"
    )
    print(
        f"Multi-label Any-Match Rate (>=1/2):       "
        f"{any_match_multi:.2%}"
    )
    print(
        f"Evidence-level Any-Match Rate:            "
        f"{evidence_any_match:.2%}"
    )

    print("-" * 72)
    print("MULTI-LABEL LABEL-LEVEL METRICS (MICRO)")

    print(f"True Positives (TP):                      {tp}")
    print(f"False Positives (FP):                     {fp}")
    print(f"False Negatives (FN):                     {fn}")
    print(f"Micro Precision:                          {micro_precision:.2%}")
    print(f"Micro Recall:                             {micro_recall:.2%}")
    print(f"Micro F1-score:                           {micro_f1:.2%}")

    print("=" * 72)


    # ========================================================
    # SAVE RESULTS
    # ========================================================

    # Prepare a human-readable ground-truth column for Excel.
    # Sets are retained internally during evaluation but are
    # exported without Python braces.
    df_merged["GT_LABELS"] = df_merged["GT_SET"].apply(
        lambda labels: " | ".join(sorted(labels))
    )

    df_merged.drop(
        columns=["GT_SET"],
        inplace=True,
    )

    output_file = OUTPUT_FILE

    df_merged.to_excel(
        output_file,
        index=False,
    )

    wb = openpyxl.load_workbook(output_file)
    ws = wb.active


    # ========================================================
    # EMBEDDED REPORT
    # ========================================================

    report_data = [
        ["METRIC", "VALUE"],
        ["TOTAL EVIDENCE ITEMS", totale],
        ["SINGLE-LABEL EVIDENCE ITEMS", totale_singoli],
        ["MULTI-LABEL EVIDENCE ITEMS", totale_doppi],
        ["SINGLE-LABEL MATCH (1/1)", match_singolo],
        ["COMPLETE MULTI-LABEL MATCH (2/2)", match_perfetto],
        ["PARTIAL MULTI-LABEL MATCH (1/2)", match_parziale],
        ["NO MATCH", nessun_match],
        ["MODEL ERRORS", errori_modello],
        ["Single-label Match Rate", f"{match_rate_singolo:.2%}"],
        [
            "Complete Multi-label Match Rate (2/2)",
            f"{complete_match_rate_multi:.2%}",
        ],
        [
            "Partial Multi-label Match Rate (1/2)",
            f"{partial_match_rate_multi:.2%}",
        ],
        [
            "Multi-label Any-Match Rate (>=1/2)",
            f"{any_match_multi:.2%}",
        ],
        [
            "Evidence-level Any-Match Rate",
            f"{evidence_any_match:.2%}",
        ],
        ["True Positives (TP)", tp],
        ["False Positives (FP)", fp],
        ["False Negatives (FN)", fn],
        ["Micro Precision", f"{micro_precision:.2%}"],
        ["Micro Recall", f"{micro_recall:.2%}"],
        ["Micro F1-score", f"{micro_f1:.2%}"],
    ]

    # Write the summary report to columns J and K.
    for r_idx, row in enumerate(report_data, 1):
        ws.cell(
            row=r_idx,
            column=10,
            value=row[0],
        )
        ws.cell(
            row=r_idx,
            column=11,
            value=row[1],
        )

    ws.cell(
        row=1,
        column=10,
    ).font = Font(bold=True)

    ws.cell(
        row=1,
        column=11,
    ).font = Font(bold=True)

    wb.save(output_file)

    print(
        f"\nEvaluation results saved to "
        f"'{output_file.name}'."
    )


else:
    print(
        "Error: model predictions and manual ground truth could "
        "not be matched. Check that the EVIDENZA texts are "
        "identical after normalization."
    )